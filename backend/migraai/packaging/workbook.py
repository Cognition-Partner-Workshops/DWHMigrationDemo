"""Deterministic XLSX and CSV exports for a packaged migration run."""

from __future__ import annotations

import csv
import io
import re
from collections import OrderedDict
from datetime import UTC, datetime
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.writer.excel import ExcelWriter

from migraai.packaging.model import PackagedObject, PackageInput

type CellValue = str | float | int | None
type SheetRows = tuple[tuple[str, ...], tuple[tuple[CellValue, ...], ...]]

CONVERTED_COLOR = "C6EFCE"
EXPERT_REVIEW_COLOR = "FFEB9C"
FAILED_COLOR = "FFC7CE"

_OBJECT_HEADERS = (
    "Source name",
    "Target name",
    "Type",
    "Source file",
    "Status",
    "Confidence",
    "Agreement",
    "Pass rate",
    "Escalated",
    "Syntax",
    "Lint",
    "Semantic",
    "Rule",
    "Dependency",
    "Object reference",
    "Security",
    "Output files",
    "Notes",
)
_CHECK_COLUMNS = (
    "syntax",
    "lint",
    "semantic",
    "rule",
    "dependency",
    "object_reference",
    "security",
)
_MAPPING_SHEETS = OrderedDict(
    (
        ("data_type", ("Data Types", "data_types.csv")),
        ("function", ("Functions", "functions.csv")),
        ("syntax", ("Syntax", "syntax.csv")),
        ("security", ("Security", "security.csv")),
    )
)
_MAPPING_HEADERS = ("Source", "Target", "Note", "KB ref", "Used by")


def _formula_escape(value: CellValue) -> CellValue:
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def _safe_relative_path(value: str) -> str:
    if value.startswith(("/", "\\")) or re.match(r"^[A-Za-z]:[\\/]", value):
        return re.split(r"[\\/]", value)[-1]
    return value


def _object_notes(obj: PackagedObject, threshold: float) -> str | None:
    parts: list[str] = []
    if obj.error:
        parts.append(obj.error)
    if obj.status == "expert_review":
        findings = [
            finding
            for check in obj.checks
            if check.status == "failed"
            for finding in check.findings
        ]
        reason = "; ".join(f"{finding.rule_id}: {finding.message}" for finding in findings)
        if not reason and obj.confidence is not None and obj.confidence < threshold:
            reason = f"confidence {obj.confidence:.2f} below threshold {threshold:.2f}"
        if not reason:
            reason = "expert review required"
        parts.append(f"NEEDS-REVIEW: {reason}")
    return "; ".join(parts) or None


def _summary_rows(pkg: PackageInput) -> tuple[tuple[str, ...], tuple[tuple[CellValue, ...], ...]]:
    counts = {
        status: sum(obj.status == status for obj in pkg.objects)
        for status in (
            "converted",
            "expert_review",
            "failed",
            "cancelled",
        )
    }
    fields: tuple[tuple[str, CellValue], ...] = (
        ("Run ID", pkg.run_id),
        ("Package name", pkg.package_name),
        ("Source platform", pkg.source_platform),
        ("Target platform", pkg.target_platform),
        ("Migration", f"{pkg.source_platform} → {pkg.target_platform}"),
        ("Created at", pkg.created_at),
        ("Confidence threshold", pkg.confidence_threshold),
        ("KB version", pkg.kb_version),
        ("Objects total", len(pkg.objects)),
        ("Converted", counts["converted"]),
        ("Expert review", counts["expert_review"]),
        ("Failed", counts["failed"]),
        ("Cancelled", counts["cancelled"]),
    )
    return ("Field", "Value"), tuple(
        (_formula_escape(label), _formula_escape(value)) for label, value in fields
    )


def _objects_rows(
    pkg: PackageInput,
) -> tuple[tuple[str, ...], tuple[tuple[CellValue, ...], ...]]:
    rows: list[tuple[CellValue, ...]] = []
    for obj in pkg.objects:
        checks = {check.check: check.status for check in obj.checks}
        row: tuple[CellValue, ...] = (
            obj.source_qualified_name,
            obj.target_qualified_name,
            obj.object_type,
            _safe_relative_path(obj.source_file),
            obj.status,
            obj.confidence,
            obj.agreement,
            obj.pass_rate,
            "yes" if obj.escalated else "no",
            *(checks.get(name) for name in _CHECK_COLUMNS),
            "; ".join(_safe_relative_path(file.package_path) for file in obj.files),
            _object_notes(obj, pkg.confidence_threshold),
        )
        rows.append(tuple(_formula_escape(value) for value in row))
    return _OBJECT_HEADERS, tuple(rows)


def _mapping_rows(
    pkg: PackageInput, kind: str
) -> tuple[tuple[str, ...], tuple[tuple[CellValue, ...], ...]]:
    mappings: dict[tuple[str, str, str, str | None], list[str]] = {}
    for obj in pkg.objects:
        for mapping in obj.mappings:
            if mapping.kind != kind:
                continue
            key = (mapping.source, mapping.target, mapping.note, mapping.kb_ref)
            used_by = mappings.setdefault(key, [])
            if obj.source_qualified_name not in used_by:
                used_by.append(obj.source_qualified_name)

    sorted_keys = sorted(
        mappings,
        key=lambda key: (key[0].casefold(), key[0], key[1], key[2], key[3] or ""),
    )
    rows: list[tuple[CellValue, ...]] = []
    for source, target, note, kb_ref in sorted_keys:
        row: tuple[CellValue, ...] = (
            source,
            target,
            note or None,
            kb_ref,
            "; ".join(mappings[(source, target, note, kb_ref)]),
        )
        rows.append(tuple(_formula_escape(value) for value in row))
    return _MAPPING_HEADERS, tuple(rows)


def _build_rows(pkg: PackageInput) -> OrderedDict[str, SheetRows]:
    sheets: OrderedDict[str, SheetRows] = OrderedDict()
    sheets["Summary"] = _summary_rows(pkg)
    sheets["Objects"] = _objects_rows(pkg)
    for kind, (sheet_name, _) in _MAPPING_SHEETS.items():
        sheets[sheet_name] = _mapping_rows(pkg, kind)
    return sheets


def _created_at(pkg: PackageInput) -> datetime:
    parsed = datetime.fromisoformat(pkg.created_at.replace("Z", "+00:00"))
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=UTC)
    return parsed.astimezone(UTC).replace(tzinfo=None)


def _zip_datetime(created_at: datetime) -> tuple[int, int, int, int, int, int]:
    return (
        min(max(created_at.year, 1980), 2107),
        created_at.month,
        created_at.day,
        created_at.hour,
        created_at.minute,
        created_at.second,
    )


def _format_sheet(
    worksheet,
    headers: tuple[str, ...],
    rows: tuple[tuple[CellValue, ...], ...],
    *,
    is_tabular: bool,
) -> None:
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)

    for row in worksheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"

    for cell in worksheet[1]:
        cell.font = Font(bold=True)
    worksheet.freeze_panes = "A2"
    if is_tabular:
        last_column = worksheet.cell(1, len(headers)).column_letter
        worksheet.auto_filter.ref = f"A1:{last_column}{max(1, worksheet.max_row)}"

    for column_index, _ in enumerate(headers, start=1):
        max_length = max(
            (
                len(str(worksheet.cell(row_index, column_index).value or ""))
                for row_index in range(1, worksheet.max_row + 1)
            ),
            default=0,
        )
        worksheet.column_dimensions[worksheet.cell(1, column_index).column_letter].width = (
            min(60, max(10, max_length)) + 2
        )


def _workbook_bytes(pkg: PackageInput, sheets: OrderedDict[str, SheetRows]) -> bytes:
    created_at = _created_at(pkg)
    workbook = Workbook()
    workbook.properties.creator = "MigraAI"
    workbook.properties.lastModifiedBy = "MigraAI"
    workbook.properties.created = created_at
    workbook.properties.modified = created_at
    workbook.properties.title = f"MigraAI mapping workbook - {pkg.package_name}"

    for sheet_index, (name, (headers, rows)) in enumerate(sheets.items()):
        worksheet = workbook.active if sheet_index == 0 else workbook.create_sheet()
        worksheet.title = name
        _format_sheet(worksheet, headers, rows, is_tabular=name != "Summary")

        if name == "Objects":
            for row_index in range(2, worksheet.max_row + 1):
                for column_index in (6, 7, 8):
                    worksheet.cell(row_index, column_index).number_format = "0.00"
            last_status_row = max(2, worksheet.max_row)
            status_range = f"E2:E{last_status_row}"
            for status, color in (
                ("converted", CONVERTED_COLOR),
                ("expert_review", EXPERT_REVIEW_COLOR),
                ("failed", FAILED_COLOR),
            ):
                worksheet.conditional_formatting.add(
                    status_range,
                    CellIsRule(
                        operator="equal",
                        formula=[f'"{status}"'],
                        fill=PatternFill(fill_type="solid", fgColor=color),
                    ),
                )

    source = io.BytesIO()
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        ExcelWriter(workbook, archive).save()

    normalized = io.BytesIO()
    timestamp = _zip_datetime(created_at)
    with ZipFile(source, "r") as original, ZipFile(normalized, "w", ZIP_DEFLATED) as archive:
        for entry in original.infolist():
            info = ZipInfo(entry.filename, date_time=timestamp)
            info.create_system = 3
            info.compress_type = ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            archive.writestr(info, original.read(entry.filename))
    return normalized.getvalue()


def build_workbook(pkg: PackageInput) -> bytes:
    """Build a deterministic XLSX mapping workbook from a packaged run."""
    return _workbook_bytes(pkg, _build_rows(pkg))


def workbook_csvs(pkg: PackageInput) -> dict[str, bytes]:
    """Build CSV exports using the same already-escaped rows as the workbook."""
    sheets = _build_rows(pkg)
    exports: dict[str, bytes] = {}
    for sheet_name, filename in (
        ("Objects", "objects.csv"),
        *((sheet_name, filename) for sheet_name, filename in _MAPPING_SHEETS.values()),
    ):
        headers, rows = sheets[sheet_name]
        output = io.StringIO(newline="")
        writer = csv.writer(output, lineterminator="\r\n")
        writer.writerow(headers)
        for row in rows:
            writer.writerow("" if value is None else value for value in row)
        exports[filename] = output.getvalue().encode("utf-8-sig")
    return exports
