from __future__ import annotations

import csv
from dataclasses import replace
from io import BytesIO
from zipfile import ZipFile

from openpyxl import load_workbook

from migraai.packaging.model import MappingRef
from migraai.packaging.sample import sample_package
from migraai.packaging.workbook import build_workbook, workbook_csvs


def _load_workbook(package=None):
    return load_workbook(BytesIO(build_workbook(package or sample_package())))


def _csv_rows(content: bytes) -> list[list[str]]:
    return list(csv.reader(content.decode("utf-8-sig").splitlines()))


def test_workbook_sheets_rows_and_summary_counts() -> None:
    package = sample_package()
    workbook = _load_workbook(package)

    assert workbook.sheetnames == [
        "Summary",
        "Objects",
        "Data Types",
        "Functions",
        "Syntax",
        "Security",
    ]
    objects = workbook["Objects"]
    assert objects.max_row == 5
    assert [objects.cell(row, 1).value for row in range(2, 6)] == [
        "sales.customer_accounts",
        "sales.v_active",
        "SELECT ON sales.v_active",
        "load_script.sql",
    ]
    assert workbook["Data Types"].max_row == 2
    assert workbook["Functions"].max_row == 2
    assert workbook["Syntax"].max_row == 2
    assert workbook["Security"].max_row == 2

    summary = {
        workbook["Summary"].cell(row, 1).value: workbook["Summary"].cell(row, 2).value
        for row in range(2, workbook["Summary"].max_row + 1)
    }
    assert summary["Converted"] == 2
    assert summary["Expert review"] == 1
    assert summary["Failed"] == 1
    assert summary["Cancelled"] == 0
    assert "All three voter sessions failed." in objects["R5"].value
    assert objects["R3"].value.startswith("NEEDS-REVIEW:")
    assert "snowflake.rule.qualify" in objects["R3"].value
    assert [objects.cell(5, column).value for column in range(10, 17)] == [None] * 7


def test_mapping_rows_deduplicate_and_keep_distinct_notes() -> None:
    package = sample_package()
    first, second, third, fourth = package.objects
    duplicate = MappingRef("data_type", "INTEGER", "NUMBER(10,0)", kb_ref="KB-1201")
    distinct_note = MappingRef(
        "data_type", "INTEGER", "NUMBER(10,0)", note="special case", kb_ref="KB-1201"
    )
    changed_second = replace(second, mappings=(duplicate,))
    changed_third = replace(third, mappings=(distinct_note,))
    package = replace(package, objects=(first, changed_second, changed_third, fourth))
    sheet = _load_workbook(package)["Data Types"]

    assert sheet.max_row == 3
    assert [sheet.cell(row, 3).value for row in range(2, 4)] == [None, "special case"]
    assert sheet["E2"].value == "sales.customer_accounts; sales.v_active"
    assert sheet["E3"].value == "SELECT ON sales.v_active"


def test_formula_injection_is_escaped_in_xlsx_and_csv() -> None:
    package = sample_package()
    obj, *remaining = package.objects
    mapping = MappingRef(
        "data_type",
        "+1",
        "-x",
        note='=HYPERLINK("http://x","y")',
    )
    obj = replace(
        obj,
        source_qualified_name="@table",
        mappings=(*obj.mappings, mapping),
    )
    package = replace(package, objects=(obj, *remaining))

    workbook = _load_workbook(package)
    objects = workbook["Objects"]
    mapping_sheet = workbook["Data Types"]
    assert objects["A2"].value == "'@table"
    assert objects["A2"].data_type == "s"
    injected = next(
        row
        for row in range(2, mapping_sheet.max_row + 1)
        if mapping_sheet.cell(row, 1).value == "'+1"
    )
    escaped_cells = tuple(mapping_sheet.cell(injected, column) for column in range(1, 6))
    assert [cell.value for cell in escaped_cells] == [
        "'+1",
        "'-x",
        '\'=HYPERLINK("http://x","y")',
        None,
        "'@table",
    ]
    assert all(escaped_cells[column].data_type == "s" for column in (0, 1, 2, 4))
    normal_mapping = next(
        row
        for row in range(2, mapping_sheet.max_row + 1)
        if mapping_sheet.cell(row, 1).value == "INTEGER"
    )
    assert mapping_sheet.cell(normal_mapping, 1).data_type == "s"

    csvs = workbook_csvs(package)
    mapping_rows = _csv_rows(csvs["data_types.csv"])
    csv_mapping = next(row for row in mapping_rows[1:] if row[0] == "'+1")
    assert csv_mapping == [
        "'+1",
        "'-x",
        '\'=HYPERLINK("http://x","y")',
        "",
        "'@table",
    ]
    assert _csv_rows(csvs["objects.csv"])[1][0] == "'@table"


def test_workbook_and_csvs_are_deterministic_and_timestamped() -> None:
    package = sample_package()
    first = build_workbook(package)
    second = build_workbook(package)
    first_csvs = workbook_csvs(package)
    second_csvs = workbook_csvs(package)
    assert first == second
    assert first_csvs == second_csvs

    with ZipFile(BytesIO(first)) as archive:
        assert all(entry.date_time == (2026, 10, 16, 0, 0, 0) for entry in archive.infolist())
        core = archive.read("docProps/core.xml").decode("utf-8")
    assert core.count("2026-10-16T00:00:00Z") == 2
    assert "<dc:creator" in core and ">MigraAI</dc:creator>" in core

    changed = replace(package, created_at="2026-10-17T00:00:00Z")
    changed_bytes = build_workbook(changed)
    assert changed_bytes != first
    with ZipFile(BytesIO(changed_bytes)) as archive:
        assert all(entry.date_time == (2026, 10, 17, 0, 0, 0) for entry in archive.infolist())


def test_csv_rows_match_workbook_sheets() -> None:
    package = sample_package()
    workbook = _load_workbook(package)
    exports = workbook_csvs(package)
    sheet_names = {
        "objects.csv": "Objects",
        "data_types.csv": "Data Types",
        "functions.csv": "Functions",
        "syntax.csv": "Syntax",
        "security.csv": "Security",
    }

    for filename, sheet_name in sheet_names.items():
        content = exports[filename]
        assert content.startswith(b"\xef\xbb\xbf")
        assert b"\r\n" in content
        assert b"\n" not in content.replace(b"\r\n", b"")
        csv_rows = _csv_rows(content)
        sheet = workbook[sheet_name]
        float_columns = {
            index
            for index, cell in enumerate(sheet[1], start=1)
            if cell.value in {"Confidence", "Agreement", "Pass rate"}
        }
        sheet_rows = [
            [
                (
                    ""
                    if sheet.cell(row, column).value is None
                    else str(float(sheet.cell(row, column).value))
                    if column in float_columns and isinstance(sheet.cell(row, column).value, int)
                    else str(sheet.cell(row, column).value)
                )
                for column in range(1, sheet.max_column + 1)
            ]
            for row in range(1, sheet.max_row + 1)
        ]
        assert csv_rows == sheet_rows


def test_workbook_formatting_and_object_status_rules() -> None:
    workbook = _load_workbook()
    assert workbook.sheetnames[0] == "Summary"
    for sheet in workbook.worksheets:
        assert sheet.freeze_panes == "A2"
        assert all(cell.font.bold for cell in sheet[1])
        if sheet.title == "Summary":
            assert sheet.auto_filter.ref is None
        else:
            assert sheet.auto_filter.ref is not None

    objects = workbook["Objects"]
    conditional_formatting = objects.conditional_formatting
    assert len(conditional_formatting) == 1
    rules = conditional_formatting["E2:E5"]
    assert len(rules) == 3
    fills = {
        rule.dxf.fill.fgColor.rgb[-6:]
        for rule in rules
        if rule.dxf is not None and rule.dxf.fill is not None
    }
    assert fills == {"C6EFCE", "FFEB9C", "FFC7CE"}
    assert {rule.operator for rule in rules} == {"equal"}
    assert {rule.formula[0] for rule in rules} == {
        '"converted"',
        '"expert_review"',
        '"failed"',
    }


def test_absolute_paths_are_reduced_and_file_content_is_not_exported() -> None:
    package = sample_package()
    obj, *remaining = package.objects
    safe_file = replace(
        obj.files[0],
        package_path=r"C:\Users\u\out.sql",
        content="CREATE TABLE PRIVATE_DATA;\nPRIVATE_CONTENT_SENTINEL",
    )
    obj = replace(obj, source_file="/home/u/x/tables.sql", files=(safe_file,))
    package = replace(package, objects=(obj, *remaining))

    workbook = _load_workbook(package)
    objects = workbook["Objects"]
    assert objects["D2"].value == "tables.sql"
    assert objects["Q2"].value == "out.sql"
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                assert "PRIVATE_CONTENT_SENTINEL" not in str(cell.value)
                assert "devin.example" not in str(cell.value)
    for content in workbook_csvs(package).values():
        decoded = content.decode("utf-8-sig")
        assert "PRIVATE_CONTENT_SENTINEL" not in decoded
        assert "devin.example" not in decoded
