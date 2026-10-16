from __future__ import annotations

import json
import uuid
import zipfile
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

from migraai.agents.job import ConversionJob
from migraai.core.audit import AuditLog
from migraai.core.crypto import EncryptedFileStore, SessionCipher
from migraai.devin.mock import MockDevinClient, MockScenario
from migraai.discovery.detect import decode_source
from migraai.discovery.inventory import Inventory
from migraai.discovery.pipeline import discover_files
from migraai.discovery.uploads import UploadRow
from migraai.packaging.naming import default_package_name
from migraai.packaging.snapshot import _upload_paths, build_package_input
from migraai.packaging.workbook import build_workbook
from migraai.packaging.zipper import build_zip
from migraai.platforms.pack import load_pack
from migraai.reports.session_report import _md, build_report_md

ROOT = Path(__file__).resolve().parents[4]
TERADATA_TESTPACK = ROOT / "testpacks" / "teradata"


def _session(directory: Path, files: dict[str, bytes]) -> dict:
    inventory = Inventory()
    store = EncryptedFileStore(directory / "store", SessionCipher())
    uploads = {}
    previews = {}
    for rel_path, data in files.items():
        row = UploadRow(
            id=uuid.uuid4().hex,
            rel_path=rel_path,
            size=len(data),
            detected_type=Path(rel_path).suffix.lstrip("."),
            status="done",
            store_path=f"uploads/{uuid.uuid4().hex}",
        )
        store.write(row.store_path, data)
        inventory.add_upload(row)
        uploads[row.id] = row
        previews[rel_path] = decode_source(data)[0]
    objects, edges = discover_files(files, "teradata")
    inventory.replace_discovery(objects, edges, uploads, previews)
    return {
        "session_id": uuid.uuid4().hex,
        "source": "teradata",
        "target": "snowflake",
        "inventory": inventory,
        "store": store,
        "uploads": uploads,
    }


def _run(
    session: dict,
    directory: Path,
    scenario: MockScenario | None = None,
    *,
    object_ids: list[str] | None = None,
    base: ConversionJob | None = None,
    replayed: bool = False,
) -> ConversionJob:
    job = ConversionJob(
        session=session,
        audit=AuditLog(directory / "audit.jsonl"),
        batch_size=20,
        timeout_minutes=2,
        max_concurrent_sessions=4,
        client_factory=lambda _: MockDevinClient(scenario),
        pack=load_pack("snowflake"),
        devin_host="local mock",
        object_ids=object_ids,
        base=base,
        replayed=replayed,
    )
    job.start()
    job.join(timeout=180)
    assert not job.is_alive()
    return job


def test_teradata_snowflake_mock_snapshot_has_56_included_objects(tmp_path: Path) -> None:
    files = {
        path.relative_to(TERADATA_TESTPACK).as_posix(): path.read_bytes()
        for path in TERADATA_TESTPACK.rglob("*.sql")
    }
    session = _session(tmp_path / "teradata", files)
    job = _run(session, tmp_path / "teradata")

    package = build_package_input(job, session, "")

    assert package.source_platform == "teradata"
    assert package.target_platform == "snowflake"
    assert package.run_id == job.run_id
    assert len(package.included()) == 56
    assert len(package.objects) == 56
    assert package.egress is not None
    assert package.egress.objects_sent == 56
    assert package.egress.devin_host == "local mock"
    assert package.objects[0].object_id == str(job.batches[0].objects[0]["id"])


def test_snapshot_propagates_kb_version_to_report_and_workbook(tmp_path: Path) -> None:
    session = _session(
        tmp_path / "kb-version",
        {"a.sql": b"CREATE TABLE sales.a (id INTEGER);"},
    )
    session["kb_version"] = 7
    job = _run(session, tmp_path / "kb-version")

    package = build_package_input(job, session, "")

    assert package.kb_version == "7"
    assert "- KB version: 7" in build_report_md(package)
    workbook = load_workbook(BytesIO(build_workbook(package)))
    summary = {
        workbook["Summary"].cell(row, 1).value: workbook["Summary"].cell(row, 2).value
        for row in range(2, workbook["Summary"].max_row + 1)
    }
    assert summary["KB version"] == "7"


def test_selected_run_merges_objects_sessions_report_and_manifest(tmp_path: Path) -> None:
    session = _session(
        tmp_path / "merge",
        {
            "a.sql": b"CREATE TABLE sales.a (id INTEGER);",
            "b.sql": b"CREATE TABLE sales.b (id INTEGER);",
        },
    )
    base = _run(session, tmp_path / "merge")
    selected_id = str(base.objects[0]["id"])
    carried_id = str(base.objects[1]["id"])
    base.results[selected_id]["confidence"] = 0.314159
    base.results[carried_id]["confidence"] = 0.271828
    base.started_sessions["cancelled-base-session"] = {
        "role": "A",
        "batch_index": 0,
        "state": "cancelled",
        "url": None,
        "acus": None,
    }

    merged = _run(
        session,
        tmp_path / "merge",
        object_ids=[selected_id],
        base=base,
    )
    package = build_package_input(merged, session, "")
    objects = {item.object_id: item for item in package.objects}

    assert len(package.objects) == 2
    assert objects[carried_id].confidence == 0.271828
    assert objects[selected_id].confidence == merged.results[selected_id]["confidence"]
    assert package.base_run_id == base.run_id
    assert package.reconverted_object_ids == (selected_id,)
    assert package.kept_base_object_ids == {}
    assert package.egress is not None
    assert package.egress.objects_sent == 3
    assert package.egress.batches == len(base.batches) + len(merged.batches)
    expected_sessions = set(base.started_sessions) | set(merged.started_sessions)
    assert package.egress.sessions == len(expected_sessions)
    assert any(
        item.run_id == base.run_id and item.state == "cancelled" for item in package.sessions
    )

    expected_note = (
        f"Re-converted 1 selected objects in run {merged.run_id}; "
        f"other results carried over from run {base.run_id}."
    )
    assert expected_note in build_report_md(package)
    with zipfile.ZipFile(BytesIO(build_zip(package, {}))) as archive:
        manifest = json.loads(archive.read(f"{package.package_name}-converted/manifest.json"))
    assert manifest["base_run_id"] == base.run_id
    assert manifest["reconverted_object_ids"] == [selected_id]
    assert manifest["kept_base_object_ids"] == {}
    assert len(manifest["objects"]) == 2
    assert "kept their earlier result" not in build_report_md(package)


def test_cancelled_selected_run_keeps_three_base_results_and_manifest(tmp_path: Path) -> None:
    files = {
        path.relative_to(TERADATA_TESTPACK).as_posix(): path.read_bytes()
        for path in TERADATA_TESTPACK.rglob("*.sql")
    }
    session = _session(tmp_path / "cancelled-merge", files)
    base = _run(session, tmp_path / "cancelled-merge")
    selected_ids = [str(item["id"]) for item in base.objects if item["type"] == "grant"][:3]
    assert len(selected_ids) == 3
    assert len(build_package_input(base, session, "").included()) == 56

    cancelled = ConversionJob(
        session=session,
        audit=AuditLog(tmp_path / "cancelled-merge" / "cancelled-audit.jsonl"),
        batch_size=20,
        timeout_minutes=2,
        client_factory=lambda _: MockDevinClient(),
        pack=load_pack("snowflake"),
        devin_host="local mock",
        object_ids=selected_ids,
        base=base,
    )
    cancelled.cancel()
    cancelled.run()

    assert cancelled.counts == base.counts
    assert cancelled.kept_base_object_ids == dict.fromkeys(selected_ids, "cancelled")
    for object_id in selected_ids:
        assert cancelled.results[object_id] == base.results[object_id]
        assert cancelled.results[object_id] is not base.results[object_id]

    package = build_package_input(cancelled, session, "")
    assert len(package.included()) == 56
    assert dict(package.kept_base_object_ids) == cancelled.kept_base_object_ids
    report = build_report_md(package)
    kept_names = sorted(
        item["qualified_name"] for item in base.objects if str(item["id"]) in selected_ids
    )
    assert (
        "3 of them kept their earlier result because the re-conversion was cancelled or "
        f"produced no usable candidate: {', '.join(_md(name) for name in kept_names)}"
    ) in report
    with zipfile.ZipFile(BytesIO(build_zip(package, {}))) as archive:
        manifest = json.loads(archive.read(f"{package.package_name}-converted/manifest.json"))
    assert manifest["kept_base_object_ids"] == cancelled.kept_base_object_ids


def test_unusable_selected_result_keeps_base_and_replaces_other_results(
    tmp_path: Path,
) -> None:
    session = _session(
        tmp_path / "unusable-merge",
        {
            "a.sql": b"CREATE TABLE sales.a (id INTEGER);",
            "b.sql": b"CREATE TABLE sales.b (id INTEGER);",
            "c.sql": b"CREATE TABLE sales.c (id INTEGER);",
        },
    )
    base = _run(session, tmp_path / "unusable-merge")
    selected_ids = [str(item["id"]) for item in base.objects]
    missing_id = selected_ids[0]
    base.results[selected_ids[1]]["confidence"] = 0.123

    reconverted = _run(
        session,
        tmp_path / "unusable-merge",
        MockScenario(missing_object_ids={"*": {missing_id}}),
        object_ids=selected_ids,
        base=base,
    )

    assert reconverted.results[missing_id] == base.results[missing_id]
    assert missing_id in reconverted.kept_base_object_ids
    assert "No usable candidate:" in reconverted.kept_base_object_ids[missing_id]
    assert reconverted.results[selected_ids[1]]["confidence"] != 0.123
    assert reconverted.results[selected_ids[2]]["confidence"] != 0.123


def test_final_packaging_failure_keeps_base_and_masks_reason(tmp_path: Path) -> None:
    session = _session(
        tmp_path / "packaging-failure",
        {"a.sql": b"CREATE TABLE sales.a (id INTEGER);"},
    )
    base = _run(session, tmp_path / "packaging-failure")
    object_id = str(base.objects[0]["id"])
    reconverted = ConversionJob(
        session=session,
        audit=AuditLog(tmp_path / "packaging-failure" / "reconvert-audit.jsonl"),
        batch_size=20,
        timeout_minutes=2,
        client_factory=lambda _: MockDevinClient(),
        pack=load_pack("snowflake"),
        devin_host="local mock",
        object_ids=[object_id],
        base=base,
    )
    store_json = reconverted._store_json

    def fail_final_file(rel_path: str, value: object) -> str:
        if rel_path.endswith("/final.json"):
            raise OSError("password=hunter2xyz")
        return store_json(rel_path, value)

    reconverted._store_json = fail_final_file
    reconverted.start()
    reconverted.join(timeout=60)
    assert not reconverted.is_alive()

    assert reconverted.results[object_id] == base.results[object_id]
    reason = reconverted.kept_base_object_ids[object_id]
    assert "***MASKED***" in reason
    assert "hunter2xyz" not in reason


def test_cancelled_reconversion_does_not_keep_failed_base_result(tmp_path: Path) -> None:
    session = _session(
        tmp_path / "failed-merge",
        {"failed.sql": b"CREATE TABLE sales.failed (id INTEGER);"},
    )
    object_id = str(session["inventory"].page(page_size=10)["items"][0]["id"])
    base = _run(
        session,
        tmp_path / "failed-merge",
        MockScenario(fail_voters={"A", "B", "C", "default"}),
    )
    assert base.results[object_id]["status"] == "failed"

    cancelled = ConversionJob(
        session=session,
        audit=AuditLog(tmp_path / "failed-merge" / "cancelled-audit.jsonl"),
        batch_size=20,
        timeout_minutes=2,
        client_factory=lambda _: MockDevinClient(),
        pack=load_pack("snowflake"),
        devin_host="local mock",
        object_ids=[object_id],
        base=base,
    )
    cancelled.cancel()
    cancelled.run()

    assert cancelled.results[object_id]["status"] == "cancelled"
    assert object_id not in cancelled.kept_base_object_ids


def test_replayed_sessions_are_labeled_and_excluded_from_egress(tmp_path: Path) -> None:
    session = _session(
        tmp_path / "replayed",
        {"a.sql": b"CREATE TABLE sales.a (id INTEGER);"},
    )
    job = _run(session, tmp_path / "replayed", replayed=True)
    package = build_package_input(job, session, "")

    assert package.egress is not None
    assert package.egress.objects_sent == 0
    assert package.egress.batches == 0
    assert package.egress.sessions == 0
    assert all(item.replayed for item in package.sessions)
    assert "replayed, nothing sent" in build_report_md(package)
    with zipfile.ZipFile(BytesIO(build_zip(package, {}))) as archive:
        manifest = json.loads(archive.read(f"{package.package_name}-converted/manifest.json"))
    assert manifest["sessions"]
    assert {item["state"] for item in manifest["sessions"]} == {"replayed, nothing sent"}


def test_failed_object_source_text_is_not_in_snapshot(tmp_path: Path) -> None:
    source = b"CREATE TABLE sales.failed_object (id INTEGER); -- source_text_marker_83726\n"
    session = _session(tmp_path / "failed", {"failed.sql": source})
    scenario = MockScenario(fail_voters={"A", "B", "C", "default"})
    job = _run(session, tmp_path / "failed", scenario)

    package = build_package_input(job, session, "")

    assert len(package.objects) == 1
    assert package.objects[0].status == "failed"
    assert package.objects[0].files == ()
    assert "source_text_marker_83726" not in repr(package)


def test_upload_roots_drive_default_package_names() -> None:
    archive = {
        str(index): SimpleNamespace(
            rel_path=f"teradata-testpack.zip/{name}.sql",
            archive_id="archive",
        )
        for index, name in enumerate(("tables", "views", "routines"))
    }
    archive_session = {"uploads": archive}
    assert _upload_paths(archive_session) == ["teradata-testpack.zip"]
    assert (
        default_package_name(_upload_paths(archive_session), "teradata", "snowflake")
        == "teradata-testpack"
    )

    mixed_session = {
        "uploads": {
            **archive,
            "loose": SimpleNamespace(rel_path="extra.sql", archive_id=None),
        }
    }
    assert _upload_paths(mixed_session) == ["teradata-testpack.zip", "extra.sql"]
    assert (
        default_package_name(_upload_paths(mixed_session), "teradata", "snowflake")
        == "migraai-teradata-to-snowflake"
    )


def test_failed_reason_is_masked_in_snapshot_report_and_manifest(tmp_path: Path) -> None:
    source = b"CREATE TABLE sales.failed_object (id INTEGER);\n"
    session = _session(tmp_path / "failed-reason", {"failed.sql": source})
    object_id = str(session["inventory"].page(page_size=300)["items"][0]["id"])
    scenario = MockScenario(
        fail_voters={"B"},
        blocked_voters={"C", "default"},
        missing_object_ids={"A": {object_id}},
    )

    class SecretErrorClient(MockDevinClient):
        def __init__(self, role: str) -> None:
            super().__init__(scenario)
            self.role = role

        def get_session(self, session_id: str) -> dict[str, object]:
            payload = super().get_session(session_id)
            if self.role == "B":
                payload["detail"] = "password=hunter2xyz\nvoter B failed"
            return payload

    job = ConversionJob(
        session=session,
        audit=AuditLog(tmp_path / "failed-reason" / "audit.jsonl"),
        batch_size=20,
        timeout_minutes=2,
        client_factory=lambda role: SecretErrorClient(role),
        pack=load_pack("snowflake"),
        devin_host="local mock",
    )
    job.start()
    job.join(timeout=60)
    assert not job.is_alive()
    reason = job.results[object_id]["reason"]
    assert "A returned no output for this object" in reason
    assert "B session failed" in reason
    assert "C session blocked" in reason
    assert "escalation session blocked" in reason
    assert "hunter2xyz" not in reason
    assert "***MASKED***" in reason
    assert job.errors[object_id] == reason

    package = build_package_input(job, session, "")
    assert package.objects[0].error == reason

    report = build_report_md(package)
    assert "No usable candidate:" in report
    assert "A returned no output for this object" in report
    assert "hunter2xyz" not in report
    parts = {"report/session-report.md": report.encode()}
    archive_bytes = build_zip(package, parts)
    with zipfile.ZipFile(BytesIO(archive_bytes)) as archive:
        manifest = json.loads(archive.read(f"{package.package_name}-converted/manifest.json"))
    manifest_reason = manifest["excluded"][0]["reason"]
    assert "hunter2xyz" not in manifest_reason
    assert "***MASKED***" in manifest_reason
